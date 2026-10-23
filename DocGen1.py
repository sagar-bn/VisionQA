import json
import os
import re
import sys
from docx import Document

# ----------------------------------------------------
# Validate Input
# ----------------------------------------------------
if len(sys.argv) < 2:
    print("Usage: py docgen.py <json_file>")
    sys.exit(1)

json_file = sys.argv[1]

# ----------------------------------------------------
# Load JSON
# ----------------------------------------------------
with open(json_file, "r", encoding="utf-8") as f:
    data = json.load(f)

# ----------------------------------------------------
# Output Folder
# ----------------------------------------------------
output_folder = "Reports"
os.makedirs(output_folder, exist_ok=True)


# ----------------------------------------------------
# Helpers
# ----------------------------------------------------
def clean_text(text):
    text = str(text)

    text = text.replace("\\r\\n", "\n")
    text = text.replace("\\n", "\n")
    text = text.replace("\\t", "    ")

    return text


def get_preview(text, max_lines=10):
    lines = text.splitlines()

    if len(lines) <= max_lines:
        return text

    return "\n".join(lines[:max_lines]) + "\n\n... truncated ..."


# ----------------------------------------------------
# Generate One DOCX Per Request
# ----------------------------------------------------
request_counter = 1

for tc in data.get("testCases", []):

    test_case_name = tc.get(
        "testCase",
        "Unknown Test Case"
    )

    requests = tc.get("requests", [])

    for request in requests:

        doc = Document()

        # ==========================================
        # Title
        # ==========================================
        doc.add_heading(
            "API Test Execution Report",
            level=1
        )

        doc.add_heading(
            test_case_name,
            level=2
        )

        # ==========================================
        # API Details
        # ==========================================
        doc.add_heading(
            "API Details",
            level=2
        )

        table = doc.add_table(
            rows=0,
            cols=2
        )

        table.style = "Table Grid"

        fields = [
            ("API Name", request.get("apiName", "")),
            ("Method", request.get("method", "")),
            ("URL", request.get("url", "")),
            ("Status", request.get("status", "")),
            ("Response Time", request.get("responseTime", "")),
            (
                "Data Size",
                f"{request.get('dataSize', '')} "
                f"{request.get('dataUnit', '')}"
            ),
        ]

        for key, value in fields:
            row = table.add_row().cells
            row[0].text = str(key)
            row[1].text = str(value)

        # ==========================================
        # Request Body
        # ==========================================
        request_body = clean_text(
            request.get("requestBody", "")
        )

        request_preview = get_preview(
            request_body,
            10
        )

        doc.add_heading(
            "Request Body Preview",
            level=2
        )

        doc.add_paragraph(
            request_preview
        )

        # Collapsible in Word
        doc.add_heading(
            "Full Request Body",
            level=2
        )

        doc.add_paragraph(
            request_body
        )

        # ==========================================
        # Response Body
        # ==========================================
        response_body = clean_text(
            request.get("responseBody", "")
        )

        response_body = response_body.replace(
            "'Response Body:',",
            ""
        )

        response_preview = get_preview(
            response_body,
            10
        )

        doc.add_heading(
            "Response Body Preview",
            level=2
        )

        doc.add_paragraph(
            response_preview
        )

        # Collapsible in Word
        doc.add_heading(
            "Full Response Body",
            level=2
        )

        doc.add_paragraph(
            response_body
        )

        # ==========================================
        # Assertions
        # ==========================================
        doc.add_heading(
            "Assertions",
            level=2
        )

        assertions = request.get(
            "assertions",
            []
        )

        if assertions:

            for assertion in assertions:

                status = assertion.get(
                    "status",
                    ""
                )

                message = assertion.get(
                    "message",
                    ""
                )

                doc.add_paragraph(
                    f"[{status.upper()}] {message}",
                    style="List Bullet"
                )

        else:
            doc.add_paragraph(
                "No Assertions Found"
            )

        # ==========================================
        # Console Logs
        # ==========================================
        doc.add_heading(
            "Console Logs",
            level=2
        )

        logs = request.get(
            "consoleLogs",
            []
        )

        if logs:

            for log in logs:
                doc.add_paragraph(
                    str(log)
                )

        else:
            doc.add_paragraph(
                "No Console Logs"
            )

        # ==========================================
        # Save DOCX
        # ==========================================
        safe_name = re.sub(
            r'[\\/*?:"<>|]',
            "_",
            test_case_name
        )

        output_file = os.path.join(
            output_folder,
            f"{request_counter}_{safe_name}.docx"
        )

        doc.save(output_file)

        print(
            f"Generated: {output_file}"
        )

        request_counter += 1

print(
    "\nAll reports generated successfully."
)
