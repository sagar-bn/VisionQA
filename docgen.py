import json
import os
import re
import sys
from docx import Document

# Check command-line argument
if len(sys.argv) < 2:
    print("Usage: py docgen.py <json_file>")
    sys.exit(1)

json_file = sys.argv[1]

# Load JSON
try:
    with open(json_file, "r", encoding="utf-8") as f:
        data = json.load(f)
except Exception as e:
    print(f"Error reading JSON file: {e}")
    sys.exit(1)

# Create Reports folder
output_folder = "Reports"
os.makedirs(output_folder, exist_ok=True)

# Verify structure
if "testCases" not in data:
    print("No 'testCases' key found in JSON.")
    sys.exit(1)

for tc in data["testCases"]:

    test_case_name = tc.get("testCase", "Unknown_TestCase")

    # Create safe filename
    safe_name = re.sub(r'[\\/*?:"<>|]', "_", test_case_name)

    output_file = os.path.join(
        output_folder,
        f"{safe_name}.docx"
    )

    doc = Document()

    # Report title
    doc.add_heading("API Test Case Report", level=1)

    doc.add_heading("Test Case", level=2)
    doc.add_paragraph(test_case_name)

    requests = tc.get("requests", [])

    doc.add_paragraph(f"Total Requests: {len(requests)}")

    for index, request in enumerate(requests, start=1):

        doc.add_page_break()

        doc.add_heading(f"Request #{index}", level=2)

        # Request Information Table
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"

        fields = [
            ("API Name", request.get("apiName", "")),
            ("Method", request.get("method", "")),
            ("URL", request.get("url", "")),
            ("Status", request.get("status", "")),
            ("Response Time", request.get("responseTime", "")),
            (
                "Data Size",
                f"{request.get('dataSize', '')} {request.get('dataUnit', '')}"
            ),
        ]

        for key, value in fields:
            row = table.add_row().cells
            row[0].text = str(key)
            row[1].text = str(value)

        # Request Body
        doc.add_heading("Request Body", level=3)

        request_body = request.get("requestBody", "")
        doc.add_paragraph(str(request_body))

        # Response Body
        doc.add_heading("Response Body", level=3)

        response_body = request.get("responseBody", "")
        doc.add_paragraph(str(response_body))

        # Assertions
        doc.add_heading("Assertions", level=3)

        assertions = request.get("assertions", [])

        if assertions:
            for assertion in assertions:
                status = assertion.get("status", "")
                message = assertion.get("message", "")

                doc.add_paragraph(
                    f"[{status.upper()}] {message}",
                    style="List Bullet"
                )
        else:
            doc.add_paragraph("No assertions found.")

        # Console Logs
        logs = request.get("consoleLogs", [])

        doc.add_heading("Console Logs", level=3)

        if logs:
            for log in logs:
                doc.add_paragraph(str(log))
        else:
            doc.add_paragraph("No console logs.")

    doc.save(output_file)

    print(f"Generated: {output_file}")

print("\nAll reports generated successfully.")